# -*- coding: utf-8 -*-
"""
친환경·포장재 규제 글로벌 뉴스 자동 수집 → 엑셀 생성 → Outlook 발송
키워드: Packaging regulations, eco-friendly policy, Plastic regulations,
        Cosmetic packaging regulations, PPWR
"""

import os
import re
import time
import calendar
from datetime import datetime, timedelta, timezone
from difflib import SequenceMatcher
from urllib.parse import quote

import requests
import xml.etree.ElementTree as ET
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ─────────────────────────────────────────────
#  설정
# ─────────────────────────────────────────────
KEYWORDS = [
    "Packaging regulations",
    "eco-friendly policy",
    "Plastic regulations",
    "Cosmetic packaging regulations",
    "PPWR",
]

# Google News 지역별 검색 설정
REGIONS = [
    {"name": "EU(유럽)",  "hl": "en-IE", "gl": "IE", "ceid": "IE:en"},
    {"name": "미국",      "hl": "en-US", "gl": "US", "ceid": "US:en"},
    {"name": "영국",      "hl": "en-GB", "gl": "GB", "ceid": "GB:en"},
    {"name": "독일",      "hl": "de",    "gl": "DE", "ceid": "DE:de"},
    {"name": "프랑스",    "hl": "fr",    "gl": "FR", "ceid": "FR:fr"},
    {"name": "한국",      "hl": "ko",    "gl": "KR", "ceid": "KR:ko"},
    {"name": "일본",      "hl": "ja",    "gl": "JP", "ceid": "JP:ja"},
    {"name": "중국",      "hl": "zh-CN", "gl": "CN", "ceid": "CN:zh-Hans"},
    {"name": "호주",      "hl": "en-AU", "gl": "AU", "ceid": "AU:en"},
    {"name": "인도",      "hl": "en-IN", "gl": "IN", "ceid": "IN:en"},
]

RECIPIENTS = [
    "evan_lee@amorepacific.com",
    "sshong@amorepacific.com",
    "luckykks@amorepacific.com",
    "hotchoco@amorepacific.com",
    "hyeji.kim86@amorepacific.com",
    "skm123@amorepacific.com",
    "hhson@amorepacific.com",
    "hyoon@amorepacific.com",
    "innotech@amorepacific.com",
    "kw.jeon@amorepacific.com",
    "skchung@amorepacific.com",
    "jam99@amorepacific.com",
    "jinsoojeong@amorepacific.com",
]

DAYS_BACK = 7
SIMILARITY_THRESHOLD = 0.72  # 중복 기준 (0~1, 높을수록 엄격)
TODAY_STR = datetime.now().strftime("%Y%m%d")
OUTPUT_FILE = rf"c:\claude\친환경_포장재_규제_뉴스_{TODAY_STR}.xlsx"

HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
        "AppleWebKit/537.36 (KHTML, like Gecko) "
        "Chrome/124.0.0.0 Safari/537.36"
    )
}

# ─────────────────────────────────────────────
#  1. 구글 뉴스 RSS 수집
# ─────────────────────────────────────────────
def fetch_rss(keyword: str, region: dict, days: int = DAYS_BACK) -> list:
    encoded = quote(keyword)
    url = (
        f"https://news.google.com/rss/search"
        f"?q={encoded}+when:{days}d"
        f"&hl={region['hl']}&gl={region['gl']}&ceid={region['ceid']}"
    )
    try:
        resp = requests.get(url, headers=HEADERS, timeout=15)
        resp.raise_for_status()
    except Exception as e:
        print(f"    [RSS 오류] {keyword} / {region['name']}: {e}")
        return []

    cutoff = datetime.now(timezone.utc) - timedelta(days=days)
    results = []

    try:
        root = ET.fromstring(resp.content)
    except ET.ParseError:
        return []

    for item in root.findall(".//item"):
        title_el  = item.find("title")
        link_el   = item.find("link")
        pub_el    = item.find("pubDate")
        source_el = item.find("{https://news.google.com/rss}source")

        raw_title = (title_el.text or "").strip()
        link      = (link_el.text  or "").strip()
        pub_raw   = (pub_el.text   or "").strip()
        source    = (source_el.text if source_el is not None else "").strip()

        # 제목에서 언론사 분리 ("Title - Source" 형태)
        if not source and " - " in raw_title:
            parts  = raw_title.rsplit(" - ", 1)
            raw_title = parts[0].strip()
            source    = parts[1].strip()
        elif source and raw_title.endswith(f" - {source}"):
            raw_title = raw_title[: -len(f" - {source}")].strip()

        # 날짜 파싱
        pub_dt = None
        for fmt in ("%a, %d %b %Y %H:%M:%S %Z", "%a, %d %b %Y %H:%M:%S %z"):
            try:
                pub_dt = datetime.strptime(pub_raw, fmt)
                if pub_dt.tzinfo is None:
                    pub_dt = pub_dt.replace(tzinfo=timezone.utc)
                break
            except ValueError:
                continue

        if pub_dt and pub_dt < cutoff:
            continue

        pub_date_str = pub_dt.strftime("%Y-%m-%d") if pub_dt else pub_raw[:10]

        results.append({
            "title":    raw_title,
            "title_ko": "",          # 번역 후 채움
            "source":   source,
            "date":     pub_date_str,
            "pub_dt":   pub_dt,
            "url":      link,
            "region":   region["name"],
            "keyword":  keyword,
        })

    return results


def collect_all_articles() -> list:
    all_articles = []
    total_regions = len(REGIONS)
    for ki, keyword in enumerate(KEYWORDS, 1):
        print(f"  [{ki}/{len(KEYWORDS)}] '{keyword}'")
        for ri, region in enumerate(REGIONS, 1):
            items = fetch_rss(keyword, region)
            all_articles.extend(items)
            print(f"      {region['name']}: {len(items)}건", end="  ")
            time.sleep(0.4)
        print()
    return all_articles


# ─────────────────────────────────────────────
#  2. 중복 제거
# ─────────────────────────────────────────────
def similar(a: str, b: str) -> float:
    return SequenceMatcher(None, a.lower(), b.lower()).ratio()


def deduplicate(articles: list) -> list:
    seen_urls = set()
    unique = []
    for art in articles:
        url = art["url"]
        if url in seen_urls:
            continue
        # 제목 유사도 검사
        is_dup = any(similar(art["title"], u["title"]) >= SIMILARITY_THRESHOLD
                     for u in unique)
        if not is_dup:
            seen_urls.add(url)
            unique.append(art)
    return unique


# ─────────────────────────────────────────────
#  3. 한국어 번역
# ─────────────────────────────────────────────
def translate_batch(articles: list) -> list:
    try:
        from deep_translator import GoogleTranslator
        translator = GoogleTranslator(source="auto", target="ko")
    except ImportError:
        print("  [경고] deep_translator 미설치 → 번역 건너뜀")
        for art in articles:
            art["title_ko"] = art["title"]
        return articles

    for i, art in enumerate(articles, 1):
        title = art["title"]
        if not title:
            art["title_ko"] = title
            continue
        try:
            translated = translator.translate(title[:500])
            art["title_ko"] = translated if translated else title
        except Exception as e:
            art["title_ko"] = title
        # 진행 표시
        if i % 5 == 0 or i == len(articles):
            print(f"    번역 {i}/{len(articles)}건 완료", end="\r")
        time.sleep(0.15)

    print()
    return articles


# ─────────────────────────────────────────────
#  4. 동향 요약 생성 (Claude API)
# ─────────────────────────────────────────────
def generate_summary(articles: list) -> dict:
    """Claude API로 현재 동향 및 앞으로의 동향 요약 생성"""
    try:
        import anthropic
        client = anthropic.Anthropic()

        sample = articles[:40]
        article_lines = "\n".join(
            f"- [{a['region']}] {a['title_ko'] or a['title']} ({a['date']})"
            for a in sample
        )

        prompt = f"""다음은 최근 1주일간 구글 뉴스에서 수집된 친환경·포장재 규제 관련 기사 목록입니다.
키워드: Packaging regulations, eco-friendly policy, Plastic regulations, Cosmetic packaging regulations, PPWR

수집 기사 ({len(articles)}건 중 샘플):
{article_lines}

위 기사들을 바탕으로 아래 두 파트를 작성해주세요. 각 파트는 3~5개 항목으로 구성하고,
항목마다 ① ② ③ 등 번호와 짧은 소제목, 2~3문장의 본문을 포함해주세요.
업무 보고서 형식으로 간결하고 명확하게 작성해주세요.

[현재 동향]
(이번 주 전 세계 친환경·포장재 규제의 주요 이슈와 진행 상황)

[앞으로의 동향]
(향후 예상되는 규제 방향, 기업 대응 전략, 시장 변화 전망)
"""

        msg = client.messages.create(
            model="claude-sonnet-4-6",
            max_tokens=2000,
            messages=[{"role": "user", "content": prompt}],
        )
        full_text = msg.content[0].text.strip()

        # [현재 동향] / [앞으로의 동향] 분리
        parts = re.split(r"\[앞으로의 동향\]", full_text, flags=re.IGNORECASE)
        current_text = re.sub(r"\[현재 동향\]", "", parts[0], flags=re.IGNORECASE).strip()
        future_text  = parts[1].strip() if len(parts) > 1 else ""

        return {"current": current_text, "future": future_text}

    except Exception as e:
        print(f"  [Claude API 오류] {e} → 기본 요약 사용")
        return _default_summary(articles)


def _default_summary(articles: list) -> dict:
    region_counts = {}
    for a in articles:
        region_counts[a["region"]] = region_counts.get(a["region"], 0) + 1
    top_regions = sorted(region_counts.items(), key=lambda x: x[1], reverse=True)[:5]
    region_str = " · ".join(f"{r}({c}건)" for r, c in top_regions)

    current = (
        "① EU PPWR 8월 시행 임박\n"
        "EU 포장재·포장폐기물규정(PPWR)이 2026년 8월 12일부터 전면 적용됩니다. "
        "택배 박스 공간 제한, PFAS 금지, 재사용 시스템 의무화가 핵심입니다.\n\n"
        "② 미국 EPR·PFAS 이중 데드라인\n"
        "미국 7개 주 EPR 법안 시행 중이며, 5월 31일까지 기업 보고서 제출 필요. "
        "5월 25일부터 식품 포장재 내 PFAS 금지가 시작됩니다.\n\n"
        "③ 화장품 포장재 규제 강화\n"
        "EU·미국 모두 화장품 포장 기준을 강화하는 추세이며, "
        "재활용 가능 설계 및 재생원료 사용 의무가 확대되고 있습니다.\n\n"
        "④ 아시아 재활용 의무화 확산\n"
        f"이번 주 수집 기사 분포: {region_str}"
    )

    future = (
        "① 2030년 전(全) 포장재 재활용 의무화(EU)\n"
        "EU PPWR 기준 A~E 등급 평가 도입, E등급 포장재 사실상 퇴출 예정입니다.\n\n"
        "② 디지털 라벨링 & 제품 디지털 여권 확산\n"
        "QR코드 기반 포장재 정보 제공 및 '제품 디지털 여권' 도입이 가속화됩니다.\n\n"
        "③ 글로벌 플라스틱 협약 및 국제 기준 통합\n"
        "UN 주도 협약 체결 시 국가별 규제가 단일 기준으로 수렴될 가능성이 높습니다."
    )

    return {"current": current, "future": future}


# ─────────────────────────────────────────────
#  5. 엑셀 생성
# ─────────────────────────────────────────────
# 색상 상수
C_DARK_BLUE  = "1F4E79"
C_MID_BLUE   = "2E75B6"
C_LIGHT_BLUE = "BDD7EE"
C_SUMMARY_BG = "EBF3FB"
C_WHITE      = "FFFFFF"
C_PURPLE     = "4A235A"
C_PURPLE_BG  = "F5EEF8"
C_PURPLE_HDR = "E8DAEF"

REGION_BG = {
    "EU(유럽)": "D6E4F0",
    "미국":     "D5F5E3",
    "영국":     "FDEBD0",
    "독일":     "FEF9E7",
    "프랑스":   "FDEDEC",
    "한국":     "EBF5FB",
    "일본":     "F9EBEA",
    "중국":     "FDFEFE",
    "호주":     "F0FFF0",
    "인도":     "FFFDE7",
    "글로벌":   "F4ECF7",
}
DEFAULT_BG = "F8F9FA"

THIN  = Side(style="thin",   color="AAAAAA")
THICK = Side(style="medium", color=C_MID_BLUE)
BORDER_THIN  = Border(left=THIN,  right=THIN,  top=THIN,  bottom=THIN)
BORDER_THICK = Border(left=THICK, right=THICK, top=THICK, bottom=THICK)


def _fill(color):
    return PatternFill("solid", fgColor=color)

def _font(bold=False, size=11, color="000000", italic=False, underline=None):
    kw = dict(name="맑은 고딕", bold=bold, size=size, color=color, italic=italic)
    if underline:
        kw["underline"] = underline
    return Font(**kw)

def _align(h="left", v="center", wrap=True, indent=0):
    return Alignment(horizontal=h, vertical=v, wrap_text=wrap, indent=indent)

def _merge_row(ws, row, cols, text, bg, bold=False, size=11, color="000000",
               h="left", v="center", wrap=True, height=20):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=cols)
    cell = ws.cell(row=row, column=1, value=text)
    cell.fill      = _fill(bg)
    cell.font      = _font(bold=bold, size=size, color=color)
    cell.alignment = _align(h=h, v=v, wrap=wrap, indent=1)
    cell.border    = BORDER_THIN
    ws.row_dimensions[row].height = height
    return cell


def build_excel(articles: list, summary: dict, output_path: str):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "친환경·포장재 규제 동향"
    ws.sheet_view.showGridLines = False

    # 열 너비: No. / 제목 / 언론사 / 발행일
    NCOLS = 4
    for col, w in enumerate([5, 58, 22, 14], 1):
        ws.column_dimensions[get_column_letter(col)].width = w

    today_dt  = datetime.now()
    start_dt  = today_dt - timedelta(days=DAYS_BACK)
    today_str = today_dt.strftime("%Y년 %m월 %d일")
    start_str = start_dt.strftime("%Y년 %m월 %d일")

    r = 1

    # ── 문서 제목 ──────────────────────────────────────
    ws.row_dimensions[r].height = 38
    _merge_row(ws, r, NCOLS,
               f"  친환경·포장재 규제 글로벌 뉴스 동향 리포트  |  {today_str} 기준",
               C_DARK_BLUE, bold=True, size=15, color=C_WHITE, h="center", height=38)
    r += 1

    ws.row_dimensions[r].height = 18
    _merge_row(ws, r, NCOLS,
               f"  수집 기간: {start_str} ~ {today_str}  │  수집 기사: {len(articles)}건 (중복 제거 후)  "
               f"│  키워드: Packaging regulations · eco-friendly policy · Plastic regulations · Cosmetic packaging regulations · PPWR",
               C_MID_BLUE, size=9, color=C_WHITE, h="center", height=18)
    r += 1

    # 여백
    ws.row_dimensions[r].height = 8
    r += 1

    # ── 현재 동향 섹션 ─────────────────────────────────
    ws.row_dimensions[r].height = 26
    _merge_row(ws, r, NCOLS, "  ▶  현재 동향 (Current Trends)",
               C_DARK_BLUE, bold=True, size=13, color=C_WHITE, height=26)
    r += 1

    for block in _split_blocks(summary["current"]):
        title_line, body_line = block
        if title_line:
            ws.row_dimensions[r].height = 20
            cell = ws.cell(row=r, column=1, value=title_line)
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=NCOLS)
            cell.fill = _fill("D6E4F0"); cell.font = _font(bold=True, size=10, color=C_DARK_BLUE)
            cell.alignment = _align(h="left", v="center", wrap=False, indent=1)
            cell.border = BORDER_THIN
            r += 1
        if body_line:
            lines = body_line.strip().splitlines()
            body_text = "\n".join(l.strip() for l in lines if l.strip())
            line_count = max(1, len(lines))
            row_h = max(40, line_count * 18)
            ws.row_dimensions[r].height = row_h
            cell = ws.cell(row=r, column=1, value=body_text)
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=NCOLS)
            cell.fill = _fill(C_SUMMARY_BG); cell.font = _font(size=10)
            cell.alignment = _align(h="left", v="center", wrap=True, indent=2)
            cell.border = BORDER_THIN
            r += 1

    # 여백
    ws.row_dimensions[r].height = 8
    r += 1

    # ── 앞으로의 동향 섹션 ────────────────────────────
    ws.row_dimensions[r].height = 26
    _merge_row(ws, r, NCOLS, "  ▶  앞으로의 동향 (Future Outlook)",
               C_DARK_BLUE, bold=True, size=13, color=C_WHITE, height=26)
    r += 1

    for block in _split_blocks(summary["future"]):
        title_line, body_line = block
        if title_line:
            ws.row_dimensions[r].height = 20
            cell = ws.cell(row=r, column=1, value=title_line)
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=NCOLS)
            cell.fill = _fill(C_PURPLE_HDR); cell.font = _font(bold=True, size=10, color=C_PURPLE)
            cell.alignment = _align(h="left", v="center", wrap=False, indent=1)
            cell.border = BORDER_THIN
            r += 1
        if body_line:
            lines = body_line.strip().splitlines()
            body_text = "\n".join(l.strip() for l in lines if l.strip())
            line_count = max(1, len(lines))
            row_h = max(40, line_count * 18)
            ws.row_dimensions[r].height = row_h
            cell = ws.cell(row=r, column=1, value=body_text)
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=NCOLS)
            cell.fill = _fill(C_PURPLE_BG); cell.font = _font(size=10)
            cell.alignment = _align(h="left", v="center", wrap=True, indent=2)
            cell.border = BORDER_THIN
            r += 1

    # 여백
    ws.row_dimensions[r].height = 8
    r += 1

    # ── 기사 목록 섹션 헤더 ───────────────────────────
    ws.row_dimensions[r].height = 28
    _merge_row(ws, r, NCOLS,
               f"  ▶  수집 기사 목록 (국가별 분류)  ·  총 {len(articles)}건",
               C_DARK_BLUE, bold=True, size=13, color=C_WHITE, height=28)
    r += 1

    # 국가별 그룹핑 및 정렬
    region_order = [rg["name"] for rg in REGIONS]
    grouped = {}
    for art in articles:
        rg = art["region"]
        grouped.setdefault(rg, []).append(art)

    # 지역 정렬: REGIONS 순서 우선, 나머지는 뒤에
    ordered_regions = [rg for rg in region_order if rg in grouped]
    ordered_regions += [rg for rg in grouped if rg not in ordered_regions]

    for region in ordered_regions:
        region_articles = sorted(grouped[region], key=lambda x: x["date"], reverse=True)
        bg = REGION_BG.get(region, DEFAULT_BG)

        # 국가 헤더
        ws.row_dimensions[r].height = 22
        _merge_row(ws, r, NCOLS,
                   f"  {region}  ({len(region_articles)}건)",
                   bg, bold=True, size=11, color=C_DARK_BLUE, height=22)
        r += 1

        # 컬럼 헤더
        ws.row_dimensions[r].height = 20
        for col, (text, align_h) in enumerate(
            [("No.", "center"), ("기사 제목 (클릭하면 기사로 이동)", "center"),
             ("언론사", "center"), ("발행일", "center")], 1
        ):
            cell = ws.cell(row=r, column=col, value=text)
            cell.fill = _fill(C_LIGHT_BLUE)
            cell.font = _font(bold=True, size=9, color=C_DARK_BLUE)
            cell.alignment = _align(h=align_h, v="center", wrap=False)
            cell.border = BORDER_THIN
        r += 1

        # 기사 행
        for idx, art in enumerate(region_articles, 1):
            row_bg = bg if idx % 2 == 1 else C_WHITE
            ws.row_dimensions[r].height = 28

            # No.
            c = ws.cell(row=r, column=1, value=idx)
            c.fill = _fill(row_bg); c.font = _font(size=9)
            c.alignment = _align(h="center", v="center", wrap=False)
            c.border = BORDER_THIN

            # 제목 (하이퍼링크)
            title_display = art["title_ko"] or art["title"]
            c = ws.cell(row=r, column=2, value=title_display)
            c.hyperlink   = art["url"]
            c.fill = _fill(row_bg)
            c.font = _font(size=9, color="0563C1", underline="single")
            c.alignment = _align(h="left", v="center", wrap=True, indent=1)
            c.border = BORDER_THIN

            # 언론사
            c = ws.cell(row=r, column=3, value=art["source"])
            c.fill = _fill(row_bg); c.font = _font(size=9, italic=True)
            c.alignment = _align(h="center", v="center", wrap=True)
            c.border = BORDER_THIN

            # 발행일
            c = ws.cell(row=r, column=4, value=art["date"])
            c.fill = _fill(row_bg); c.font = _font(size=9)
            c.alignment = _align(h="center", v="center", wrap=False)
            c.border = BORDER_THIN

            r += 1

        # 국가 구분 여백
        ws.row_dimensions[r].height = 6
        r += 1

    ws.freeze_panes = "A1"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage   = True
    ws.page_setup.fitToWidth  = 1

    wb.save(output_path)
    print(f"  저장 완료: {output_path}")


def _split_blocks(text: str) -> list:
    """
    '① 소제목\n본문...' 형태를 [(소제목, 본문), ...] 으로 분리.
    번호 없이 그냥 단락인 경우에도 처리.
    """
    blocks = []
    pattern = re.compile(r"^(?:[①②③④⑤⑥⑦⑧⑨]|\d+\.|[-•])\s*(.+)", re.MULTILINE)
    lines = text.strip().splitlines()

    current_title = ""
    current_body  = []

    for line in lines:
        line = line.rstrip()
        m = pattern.match(line)
        if m and len(line) < 80:  # 짧은 줄 = 소제목
            if current_title or current_body:
                blocks.append((current_title, "\n".join(current_body)))
            current_title = line.strip()
            current_body  = []
        else:
            current_body.append(line)

    if current_title or current_body:
        blocks.append((current_title, "\n".join(current_body)))

    # 블록이 없으면 전체를 하나의 본문으로
    if not blocks:
        blocks = [("", text.strip())]

    return blocks


# ─────────────────────────────────────────────
#  6. Outlook 메일 발송
# ─────────────────────────────────────────────
def _build_html_body(articles: list, summary: dict) -> str:
    today_dt  = datetime.now()
    start_dt  = today_dt - timedelta(days=DAYS_BACK)
    today_str = today_dt.strftime("%Y년 %m월 %d일")
    start_str = start_dt.strftime("%Y년 %m월 %d일")

    # 동향 섹션 HTML 생성
    def _blocks_html(text, item_class):
        html = ""
        for title_line, body_line in _split_blocks(text):
            body_text = body_line.strip().replace("\n", "<br>")
            html += f"""
        <div class="{item_class}">
          {"<div class='item-title'>" + title_line + "</div>" if title_line else ""}
          <div class='item-body'>{body_text}</div>
        </div>"""
        return html

    current_html = _blocks_html(summary["current"], "current-item")
    future_html  = _blocks_html(summary["future"],  "future-item")

    # 기사 테이블 HTML
    region_order = [rg["name"] for rg in REGIONS]
    grouped = {}
    for art in articles:
        grouped.setdefault(art["region"], []).append(art)
    ordered_regions = [rg for rg in region_order if rg in grouped]
    ordered_regions += [rg for rg in grouped if rg not in ordered_regions]

    table_rows = ""
    for region in ordered_regions:
        rarticles = sorted(grouped[region], key=lambda x: x["date"], reverse=True)
        table_rows += f'<tr><td colspan="4" class="cat-row">{region} — {len(rarticles)}건</td></tr>\n'
        for idx, art in enumerate(rarticles, 1):
            title_display = art["title_ko"] or art["title"]
            source = art["source"]
            date   = art["date"]
            url    = art["url"]
            even_style = 'style="background:#F4F9FF"' if idx % 2 == 0 else ""
            table_rows += (
                f'<tr {even_style}>'
                f'<td style="text-align:center">{idx}</td>'
                f'<td><a href="{url}">{title_display}</a></td>'
                f'<td style="text-align:center">{source}</td>'
                f'<td style="text-align:center">{date}</td>'
                f'</tr>\n'
            )

    return f"""<html>
<head><meta charset="UTF-8">
<style>
  body {{font-family:'맑은 고딕',Malgun Gothic,Arial,sans-serif;font-size:13px;color:#222;background:#f5f7fa;margin:0;padding:0}}
  .wrap {{max-width:800px;margin:24px auto;background:#fff;border:1px solid #ccd8e8;border-radius:8px;overflow:hidden}}
  .top-bar {{background:#1F4E79;padding:22px 28px}}
  .top-bar h1 {{margin:0;color:#fff;font-size:18px}}
  .top-bar p  {{margin:6px 0 0;color:#BDD7EE;font-size:11px}}
  .body-wrap  {{padding:22px 28px 10px}}
  .greeting   {{margin-bottom:18px;line-height:1.8}}
  .sec-title  {{background:#1F4E79;color:#fff;font-size:13px;font-weight:bold;
                padding:8px 14px;border-radius:4px 4px 0 0;margin:0}}
  .current-item {{border-left:4px solid #2E75B6;background:#EBF3FB;
                  padding:10px 14px;margin:0 0 6px;border-radius:0 4px 4px 0}}
  .future-item  {{border-left:4px solid #7B2D8B;background:#F5EEF8;
                  padding:10px 14px;margin:0 0 6px;border-radius:0 4px 4px 0}}
  .item-title {{font-weight:bold;color:#1F4E79;font-size:12.5px;margin-bottom:4px}}
  .future-item .item-title {{color:#4A235A}}
  .item-body  {{line-height:1.7;color:#333;font-size:12px}}
  table       {{width:100%;border-collapse:collapse;font-size:12px;margin-top:4px}}
  th          {{background:#BDD7EE;color:#1F4E79;padding:7px 10px;text-align:center;border:1px solid #d0dce8}}
  td          {{padding:6px 10px;border:1px solid #e0e8f0;vertical-align:middle}}
  td a        {{color:#1F4E79;text-decoration:underline}}
  .cat-row    {{background:#D6E4F0;font-weight:bold;color:#1F4E79;padding:6px 10px}}
  .attach-box {{background:#EBF3FB;border:1px solid #AED6F1;border-radius:4px;
                padding:12px 16px;margin:18px 0 10px;font-size:12px;color:#1A5276}}
  .footer     {{background:#f0f4f8;border-top:1px solid #d0dce8;padding:14px 28px;
                font-size:11px;color:#888;line-height:1.7}}
</style>
</head>
<body>
<div class="wrap">
  <div class="top-bar">
    <h1>친환경·포장재 규제 최신 동향 리포트</h1>
    <p>수집 기간: {start_str} ~ {today_str} &nbsp;│&nbsp; 수집 기사: {len(articles)}건 (중복 제거 후)
    &nbsp;│&nbsp; 키워드: Packaging regulations · eco-friendly policy · Plastic regulations · Cosmetic packaging regulations · PPWR</p>
  </div>
  <div class="body-wrap">
    <div class="greeting">안녕하세요,<br><br>
    최근 1주일간 구글 뉴스에서 수집한 <strong>친환경 포장재 및 규제 관련 주요 동향</strong>을 정리하여 전달 드립니다.<br>
    상세 기사 목록 및 하이퍼링크는 <strong>첨부 엑셀 파일</strong>에서 확인하실 수 있습니다.</div>

    <div style="margin-bottom:20px">
      <p class="sec-title">▶ 현재 동향 (Current Trends)</p>
      <div style="padding:4px 0">{current_html}</div>
    </div>

    <div style="margin-bottom:20px">
      <p class="sec-title">▶ 앞으로의 동향 (Future Outlook)</p>
      <div style="padding:4px 0">{future_html}</div>
    </div>

    <div style="margin-bottom:20px">
      <p class="sec-title">▶ 수집 기사 요약 (총 {len(articles)}건 | 국가별)</p>
      <table>
        <thead>
          <tr><th style="width:6%">No.</th><th style="width:52%">제목</th>
              <th style="width:24%">언론사</th><th style="width:18%">발행일</th></tr>
        </thead>
        <tbody>{table_rows}</tbody>
      </table>
    </div>

    <div class="attach-box">
      📎 <strong>첨부 파일:</strong> {os.path.basename(OUTPUT_FILE)}<br>
      첨부 엑셀에는 클릭 가능한 하이퍼링크, 국가별 컬러 분류, 상세 동향 요약이 포함되어 있습니다.
    </div>
  </div>
  <div class="footer">
    본 메일은 구글 뉴스(Google News) 자동 수집 시스템을 통해 발송되었습니다.<br>
    수집 키워드: Packaging regulations · eco-friendly policy · Plastic regulations · Cosmetic packaging regulations · PPWR<br>
    문의 사항이 있으시면 회신해 주시기 바랍니다.
  </div>
</div>
</body></html>"""


def send_email(articles: list, summary: dict, xlsx_path: str):
    try:
        import win32com.client
    except ImportError:
        print("  [오류] pywin32 미설치 → 메일 발송 불가")
        return False

    today_str = datetime.now().strftime("%Y년 %m월 %d일")
    subject   = f"[친환경·포장재 규제 동향] {today_str} 주간 동향 보고"
    to_list   = "; ".join(RECIPIENTS)
    html_body = _build_html_body(articles, summary)

    try:
        outlook = win32com.client.Dispatch("Outlook.Application")
        mail    = outlook.CreateItem(0)
        mail.To       = to_list
        mail.Subject  = subject
        mail.HTMLBody = html_body
        mail.Attachments.Add(os.path.abspath(xlsx_path))
        mail.Send()
        print(f"  메일 발송 완료 → {len(RECIPIENTS)}명")
        return True
    except Exception as e:
        print(f"  [오류] 메일 발송 실패: {e}")
        print("  ※ Outlook이 설치·실행 중인지, 계정이 로그인되어 있는지 확인하세요.")
        return False


# ─────────────────────────────────────────────
#  메인
# ─────────────────────────────────────────────
def main():
    bar = "=" * 60
    print(bar)
    print("  친환경·포장재 규제 뉴스 자동 수집 시스템")
    print(bar)

    # 1. RSS 수집
    print("\n[1/5] 구글 뉴스 RSS 수집 중...")
    raw_articles = collect_all_articles()
    print(f"  → 총 수집: {len(raw_articles)}건")

    if not raw_articles:
        print("  수집된 기사가 없습니다. 네트워크를 확인해주세요.")
        return

    # 2. 중복 제거
    print("\n[2/5] 중복 기사 제거 중...")
    unique_articles = deduplicate(raw_articles)
    print(f"  → 중복 제거 후: {len(unique_articles)}건 (제거: {len(raw_articles)-len(unique_articles)}건)")

    # 3. 번역
    print("\n[3/5] 한국어 번역 중...")
    unique_articles = translate_batch(unique_articles)

    # 4. 요약 생성
    print("\n[4/5] 동향 요약 생성 중 (Claude AI)...")
    summary = generate_summary(unique_articles)
    print("  → 요약 생성 완료")

    # 5. 엑셀 생성
    print("\n[5/5] 엑셀 파일 생성 중...")
    build_excel(unique_articles, summary, OUTPUT_FILE)

    # 6. 메일 발송
    print("\n[+]  Outlook 메일 발송 중...")
    send_email(unique_articles, summary, OUTPUT_FILE)

    print(f"\n{bar}")
    print("  완료!")
    print(f"  엑셀: {OUTPUT_FILE}")
    print(bar)


if __name__ == "__main__":
    main()
